"""Data sources: tables that graphics read from.

A source is a Google Sheet, a CSV / TSV / Excel / text / JSON file (added
to the engine, or a file elsewhere on this PC that is watched for changes),
or a web address giving CSV or JSON. Each is read into a table (rows of
text) and read again every few seconds.

In text, {{Name!B2}} is the cell B2 of the source called Name,
{{Name!Team#3}} the "Team" column of the 3rd row under the headers, and
{{Name!Score}} the cell next to the one saying "Score" in column A (a
two-column "label | value" sheet).

    {"id", "name", "type": "gsheet" | "file" | "url", "url" or "path",
     "sheet": "" (Excel sheet / Google Sheet tab), "refresh": seconds}
"""

import csv
import io
import json
import os
import re
import threading
import time
import urllib.parse
import urllib.request

from .log import log

TEXT_EXTS = (".csv", ".tsv", ".txt", ".json", ".xlsx", ".xlsm")
MAX_ROWS = 2000
MAX_COLS = 100


def gsheet_csv_url(url, sheet=""):
    """A Google Sheets link (edit, view, or published) -> its CSV export link."""
    if "output=csv" in url or "format=csv" in url or "tqx=out:csv" in url:
        return url
    m = re.search(r"/spreadsheets/d/(e/)?([\w-]+)", url)
    if not m:
        raise ValueError("That isn't a Google Sheets link")
    if m.group(1):   # "Publish to web" link
        gid = re.search(r"[#&?]gid=(\d+)", url)
        return f"https://docs.google.com/spreadsheets/d/e/{m.group(2)}/pub?output=csv" + (f"&gid={gid.group(1)}" if gid else "")
    if sheet:
        return (f"https://docs.google.com/spreadsheets/d/{m.group(2)}/gviz/tq?tqx=out:csv&sheet="
                + urllib.parse.quote(sheet))
    gid = re.search(r"[#&?]gid=(\d+)", url)
    return f"https://docs.google.com/spreadsheets/d/{m.group(2)}/export?format=csv" + (f"&gid={gid.group(1)}" if gid else "")


def _clip(rows):
    rows = [[("" if c is None else str(c)) for c in r[:MAX_COLS]] for r in rows[:MAX_ROWS]]
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def parse_csv(text, delimiter=None):
    if delimiter is None:
        sample = text[:4096]
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    return _clip(list(csv.reader(io.StringIO(text), delimiter=delimiter)))


def parse_json(text):
    data = json.loads(text)
    if isinstance(data, dict):
        # {"key": value} -> label | value rows; {"rows": [...]} -> those rows.
        lists = [v for v in data.values() if isinstance(v, list)]
        if len(data) == 1 and lists:
            data = lists[0]
        else:
            return _clip([[k, v if isinstance(v, str) else json.dumps(v)] for k, v in data.items()])
    if isinstance(data, list) and data and isinstance(data[0], dict):
        headers = []
        for item in data:
            for k in item:
                if k not in headers:
                    headers.append(k)
        return _clip([headers] + [[item.get(h, "") if isinstance(item.get(h, ""), str) else json.dumps(item.get(h))
                                   for h in headers] for item in data])
    if isinstance(data, list):
        return _clip([r if isinstance(r, list) else [r] for r in data])
    return [[str(data)]]


def parse_text(text):
    return _clip([[line] for line in text.splitlines()])


def parse_xlsx(blob, sheet=""):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else _cell_text(v) for v in row])
        if len(rows) >= MAX_ROWS:
            break
    sheets = wb.sheetnames
    wb.close()
    return _clip(rows), sheets


def _cell_text(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def parse(name, blob, sheet=""):
    """(rows, sheet names) for a file's contents, by its extension."""
    ext = os.path.splitext(name.lower())[1]
    if ext in (".xlsx", ".xlsm"):
        return parse_xlsx(blob, sheet)
    text = blob.decode("utf-8-sig", "replace")
    if ext == ".json":
        return parse_json(text), []
    if ext == ".txt":
        return parse_text(text), []
    if ext == ".tsv":
        return parse_csv(text, "\t"), []
    if ext == ".csv":
        return parse_csv(text), []
    stripped = text.lstrip()
    if stripped.startswith(("{", "[")):
        return parse_json(text), []
    return parse_csv(text), []


def _fetch(url):
    request = urllib.request.Request(url, headers={"User-Agent": "SCC-GFX-Engine"})
    with urllib.request.urlopen(request, timeout=15) as r:
        return r.read(), r.headers.get("Content-Type", "")


class Sources:
    """Reads every data source, now and every few seconds after."""

    def __init__(self, workspace):
        self.ws = workspace
        self.tables = {}   # id -> {"rows", "sheets", "updated", "error"}
        self.lock = threading.Lock()
        self.listeners = []
        self.rev = 0
        self._next = {}
        self._stop = threading.Event()
        threading.Thread(target=self._loop, daemon=True).start()

    def stop(self):
        self._stop.set()

    def snapshot(self):
        """{source name: rows} for the outputs, plus each table's status."""
        sources = self.ws.snapshot()["sources"]
        with self.lock:
            return {"rev": self.rev,
                    "tables": {s["name"]: self.tables.get(s["id"], {}).get("rows", []) for s in sources},
                    "status": {s["id"]: {k: v for k, v in self.tables.get(s["id"], {}).items() if k != "rows"}
                               for s in sources}}

    def table(self, sid):
        with self.lock:
            return dict(self.tables.get(sid, {}))

    def refresh(self, sid=None):
        for s in self.ws.get()["sources"]:
            if sid is None or s["id"] == sid:
                self._read(s)

    def _read(self, source):
        try:
            kind = source.get("type")
            if kind == "gsheet":
                blob, _ = _fetch(gsheet_csv_url(source["url"], source.get("sheet", "")))
                if blob.lstrip()[:15].lower().startswith((b"<!doctype", b"<html")):
                    raise ValueError("Google wants a sign-in: share the sheet as \"Anyone with the link can view\"")
                rows, sheets = parse_csv(blob.decode("utf-8-sig", "replace")), []
            elif kind == "url":
                blob, ctype = _fetch(source["url"])
                name = source["url"].split("?")[0]
                if "json" in ctype and not name.endswith(".json"):
                    name += ".json"
                rows, sheets = parse(name, blob, source.get("sheet", ""))
            else:
                path = source.get("path", "")
                if path.startswith("assets/"):
                    path = self.ws.asset_path(path) or ""
                with open(path, "rb") as f:
                    blob = f.read()
                rows, sheets = parse(path, blob, source.get("sheet", ""))
            entry = {"rows": rows, "sheets": sheets, "updated": time.time(), "error": None}
        except Exception as e:
            entry = dict(self.table(source["id"]))
            entry["error"] = str(getattr(e, "reason", e))
            entry.setdefault("rows", [])
        with self.lock:
            old = self.tables.get(source["id"], {})
            changed = old.get("rows") != entry["rows"] or old.get("error") != entry["error"]
            self.tables[source["id"]] = entry
            if changed:
                self.rev += 1
        if changed:
            if entry["error"]:
                log(f"Data source {source['name']}: {entry['error']}", "warn")
            for fn in list(self.listeners):
                fn()

    def _loop(self):
        while not self._stop.wait(0.5):
            now = time.time()
            for s in self.ws.get()["sources"]:
                every = max(1.0, float(s.get("refresh") or 5))
                if now >= self._next.get(s["id"], 0):
                    self._next[s["id"]] = now + every
                    threading.Thread(target=self._read, args=(s,), daemon=True).start()

    def forget(self, sid):
        with self.lock:
            self.tables.pop(sid, None)
            self.rev += 1
        self._next.pop(sid, None)


# --- references -------------------------------------------------------

def col_index(letters):
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def col_letters(index):
    out = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def lookup(rows, ref):
    """A cell's text by reference: "B2", "Team#3" or "Score" (see the module doc)."""
    ref = ref.strip()
    m = re.fullmatch(r"([A-Za-z]{1,3})(\d{1,5})", ref)
    if m:
        r, c = int(m.group(2)) - 1, col_index(m.group(1))
        return rows[r][c] if 0 <= r < len(rows) and 0 <= c < len(rows[r]) else ""
    if "#" in ref and rows:
        header, _, n = ref.rpartition("#")
        headers = [h.strip().lower() for h in rows[0]]
        if header.strip().lower() in headers and n.strip().isdigit():
            c, r = headers.index(header.strip().lower()), int(n)
            return rows[r][c] if r < len(rows) and c < len(rows[r]) else ""
    for row in rows:
        if row and row[0].strip().lower() == ref.lower():
            return row[1] if len(row) > 1 else ""
    return ""
